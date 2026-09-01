from pathlib import Path
import sys
import re
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"C:\Users\jongp\Desktop\Thinkpad_and_PC_Sync\ContentAutomationStudio")
version = sys.argv[1] if len(sys.argv) > 1 else "v0.1"
src = ROOT / "docs" / f"PRD_Content_Automation_Studio_{version}.md"
out = ROOT / "docs" / f"PRD_Content_Automation_Studio_{version}.docx"

lines = src.read_text(encoding="utf-8").splitlines()
doc = Document()
sec = doc.sections[0]
sec.top_margin = Mm(18)
sec.bottom_margin = Mm(18)
sec.left_margin = Mm(20)
sec.right_margin = Mm(20)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Tahoma")
styles["Normal"].font.size = Pt(10.5)
for name, size, color in [("Title", 26, "3C2A78"), ("Heading 1", 18, "3C2A78"), ("Heading 2", 14, "4F3E8C"), ("Heading 3", 12, "5F557A")]:
    st = styles[name]
    st.font.name = "Aptos Display"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Tahoma")
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor.from_string(color)
    st.font.bold = True

header = sec.header.paragraphs[0]
header.text = "CONTENT AUTOMATION STUDIO  |  PRODUCT REQUIREMENTS DOCUMENT"
header.alignment = WD_ALIGN_PARAGRAPH.CENTER
header.runs[0].font.size = Pt(8)
header.runs[0].font.color.rgb = RGBColor(110, 110, 110)

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run(f"Draft {version}  •  WordQuest Project  •  Page ")
r.font.size = Pt(8)
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
footer._p.append(fld)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def add_inline(paragraph, text):
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            paragraph.add_run(part)


def parse_table(start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        vals = [x.strip() for x in lines[i].strip().strip("|").split("|")]
        rows.append(vals)
        i += 1
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", x.replace(" ", "")) for x in rows[1]):
        rows.pop(1)
    return rows, i

# Cover
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("CONTENT AUTOMATION\nSTUDIO")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Product Requirements Document\nStoryboard • ComfyUI • Automated Video Assembly")
r.bold = True
r.font.size = Pt(16)
r.font.color.rgb = RGBColor(47, 184, 166)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Version 0.1  |  Draft for Review").italic = True
doc.add_page_break()

in_code = False
code_lines = []
i = 2  # Skip the markdown title lines already represented on cover
while i < len(lines):
    raw = lines[i]
    s = raw.strip()
    if s.startswith("```"):
        if not in_code:
            in_code = True
            code_lines = []
        else:
            p = doc.add_paragraph()
            p.style = styles["Normal"]
            r = p.add_run("\n".join(code_lines))
            r.font.name = "Consolas"
            r.font.size = Pt(8.5)
            in_code = False
        i += 1
        continue
    if in_code:
        code_lines.append(raw)
        i += 1
        continue
    if not s or s == "---":
        i += 1
        continue
    if s.startswith("|"):
        rows, i = parse_table(i)
        if rows:
            cols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=cols)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ri, row in enumerate(rows):
                for ci in range(cols):
                    cell = table.cell(ri, ci)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    cell.text = row[ci] if ci < len(row) else ""
                    for para in cell.paragraphs:
                        for run in para.runs:
                            run.font.size = Pt(9)
                            if ri == 0:
                                run.bold = True
                                run.font.color.rgb = RGBColor(255, 255, 255)
                    if ri == 0:
                        shade_cell(cell, "4F3E8C")
                    elif ri % 2 == 0:
                        shade_cell(cell, "F2F0FA")
            doc.add_paragraph()
        continue
    m = re.match(r"^(#{1,6})\s+(.*)$", s)
    if m:
        level = len(m.group(1))
        title = m.group(2)
        if level == 1:
            doc.add_heading(title, level=1)
        else:
            doc.add_heading(title, level=min(level - 1, 3))
        i += 1
        continue
    if s.startswith(">"):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Mm(8)
        p.paragraph_format.right_indent = Mm(8)
        r = p.add_run(s.lstrip("> "))
        r.italic = True
        r.font.color.rgb = RGBColor(79, 62, 140)
        i += 1
        continue
    if re.match(r"^[-*]\s+", s):
        p = doc.add_paragraph(style="List Bullet")
        add_inline(p, re.sub(r"^[-*]\s+", "", s))
        i += 1
        continue
    if re.match(r"^\d+\.\s+", s):
        p = doc.add_paragraph(style="List Number")
        add_inline(p, re.sub(r"^\d+\.\s+", "", s))
        i += 1
        continue
    p = doc.add_paragraph()
    add_inline(p, s.replace("  ", " "))
    i += 1

doc.core_properties.title = "Content Automation Studio — Product Requirements Document"
doc.core_properties.subject = "AI Content Automation, Storyboard, ComfyUI and Video Assembly PRD"
doc.core_properties.author = "Gutenberg"
doc.core_properties.keywords = "Content Automation, PRD, Storyboard, ComfyUI, Video Generation, FFmpeg"
doc.save(out)
print(out)
